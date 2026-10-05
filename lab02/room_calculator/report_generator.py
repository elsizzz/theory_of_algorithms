from docx import Document
from .models import Building

class ReportGenerator:
    """Класс для сохранения отчёта в формате .docx."""
    
    @staticmethod
    def save_to_word(building: Building, filename="report.docx"):
        doc = Document()
        doc.add_heading(f'Отчёт по зданию: {building.name}', 0)

        doc.add_heading('Общие показатели', level=1)
        doc.add_paragraph(f'Общая площадь: {building.total_area():.2f} кв.м')
        doc.add_paragraph(f'Требуемая тепловая мощность: {building.total_heating_power():.2f} кВт')

        doc.add_heading('Детализация по квартирам', level=1)
        for apt in building.apartments:
            doc.add_heading(f'Квартира: {apt.name}', level=2)
            doc.add_paragraph(f'Площадь: {apt.total_area():.2f} кв.м | Тепловая мощность: {apt.total_heating_power():.2f} кВт')
            
            for room in apt.rooms:
                doc.add_paragraph(
                    f'- {room.name}: Площадь {room.area():.2f} кв.м, '
                    f'Тепловая мощность {room.heating_power():.2f} кВт (Утепление: {room.insulation})'
                )

        doc.save(filename)
        print(f"\n[Успех] Отчёт сохранён в файл: {filename}")